import hashlib
import json
import logging
import math
import re
import time
from datetime import date, datetime
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

import yaml
from google import genai
from google.genai import errors, types
from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langsmith import traceable
from openpyxl import load_workbook

from backend.app.config import PROJECT_ROOT, get_settings
from backend.app.database import get_connection

PRODUCT_DOCUMENTS_DIR = PROJECT_ROOT / "docs" / "product"
SUPPORTED_EXTENSIONS = {".md", ".pdf", ".docx", ".xlsx"}
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 150
EMBEDDING_BATCH_SIZE = 32
logger = logging.getLogger(__name__)


def normalize_embedding_vector(values: list[float]) -> list[float]:
    squared_values: list[float] = []

    for value in values:
        squared_values.append(math.pow(value, 2))

    length = math.sqrt(sum(squared_values))

    if length == 0:
        raise RuntimeError("Embedding service returned a zero vector")

    normalized_values: list[float] = []

    for value in values:
        normalized_values.append(value / length)

    return normalized_values


@traceable(name="google_embedding", run_type="embedding")
def generate_text_embeddings(texts: list[str], task_type: str) -> list[list[float]]:
    settings = get_settings()
    client = genai.Client(api_key=settings.embedding_api_key)

    result = client.models.embed_content(
        model=settings.embedding_model,
        contents=texts,
        config=types.EmbedContentConfig(
            task_type=task_type,
            output_dimensionality=settings.embedding_dimension,
        ),
    )

    client.close()

    embedding_items = result.embeddings

    if embedding_items is None:
        embedding_items = []

    embedding_vectors: list[list[float]] = []

    for item in embedding_items:
        item_values = item.values

        if item_values is None:
            values: list[float] = []
        else:
            values = list(item_values)

        normalized_vector = normalize_embedding_vector(values)
        embedding_vectors.append(normalized_vector)

    if len(embedding_vectors) != len(texts):
        raise RuntimeError(
            "Embedding response did not match the requested count or dimension"
        )

    for embedding_vector in embedding_vectors:
        if len(embedding_vector) != settings.embedding_dimension:
            raise RuntimeError(
                "Embedding response did not match the requested count or dimension"
            )

    return embedding_vectors


def discover_product_documents() -> list[Path]:
    if not PRODUCT_DOCUMENTS_DIR.is_dir():
        raise ValueError(f"Product documents directory is missing: {PRODUCT_DOCUMENTS_DIR}")
    paths = []
    for path in PRODUCT_DOCUMENTS_DIR.rglob("*"):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            paths.append(path)
    if not paths:
        raise ValueError("No product documents found; refusing to clear the knowledge base")
    return sorted(paths)


def metadata_date(value: object, field: str, path: Path) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError as error:
        raise ValueError(f"Invalid {field}: {path}") from error


def validate_document_metadata(metadata: dict[str, object], path: Path) -> dict[str, object]:
    required_fields = {"title", "company_id", "product", "version", "effective_from", "status", "last_reviewed", "topic"}
    missing_fields = required_fields.difference(metadata)
    if missing_fields:
        raise ValueError(f"Missing metadata {sorted(missing_fields)}: {path}")
    for field in ("title", "company_id", "product", "version", "status", "topic"):
        if metadata[field] is None or not str(metadata[field]).strip():
            raise ValueError(f"Empty {field}: {path}")
        metadata[field] = str(metadata[field]).strip()
    metadata["effective_from"] = metadata_date(metadata["effective_from"], "effective_from", path)
    effective_to = metadata.get("effective_to")
    if effective_to is not None:
        effective_to = metadata_date(effective_to, "effective_to", path)
        if effective_to < metadata["effective_from"]:
            raise ValueError(f"Invalid effective date range: {path}")
    metadata["effective_to"] = effective_to
    metadata["last_reviewed"] = metadata_date(metadata["last_reviewed"], "last_reviewed", path)
    metadata["source_uri"] = path.relative_to(PROJECT_ROOT).as_posix()
    return metadata


def extract_source_metadata(content: str, path: Path, effective_dates: dict) -> dict[str, object]:
    lines = []
    for line in content.splitlines():
        if line.strip():
            lines.append(line.strip())
    metadata: dict[str, object] = {"title": lines[0].lstrip("# ") if lines else ""}
    labels = {"company_id": "公司", "product": "产品", "version": "版本", "status": "状态", "last_reviewed": "最后核对日期", "topic": "主题", "effective_from": "effective_from", "effective_to": "effective_to"}
    for field, label in labels.items():
        match = re.search(re.escape(label) + r"[：:]\s*([^|\n]+)", content)
        if match:
            metadata[field] = match.group(1).strip()
    if "effective_from" not in metadata:
        scope = (metadata.get("company_id"), metadata.get("product"), metadata.get("version"))
        periods = effective_dates.get(scope, set())
        if len(periods) != 1:
            raise ValueError(f"Missing or ambiguous effective dates for {scope}: {path}")
        metadata["effective_from"], metadata["effective_to"] = next(iter(periods))
    return metadata


def read_excel_document(path: Path) -> tuple[str, list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    records: list[dict[str, object]] = []
    try:
        for sheet in workbook:
            header = None
            intro_lines: list[str] = []
            record_count = 0
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = []
                for value in row:
                    if isinstance(value, datetime):
                        values.append(value.date().isoformat())
                    elif value is None:
                        values.append("")
                    else:
                        values.append(str(value).strip())
                if not any(values):
                    continue
                if header is None:
                    # 当前参考表以这组业务列标识记录区，前面的行属于来源信息与共同规则。
                    if "含义" in values and "用户下一步" in values and "条件 / 必要证据" in values:
                        header = values
                        intro = "\n".join(intro_lines)
                        sections.append(intro)
                        records.append({"sheet": sheet.title, "row": "概述", "content": intro})
                    else:
                        populated = [value for value in values if value]
                        intro_lines.append("：".join(populated))
                    continue
                fields = []
                for label, value in zip(header, values, strict=True):
                    if value and not label:
                        raise ValueError(f"Excel value without a header: {path}, row {row_number}")
                    if label:
                        fields.append(f"{label}：{value}")
                content = "\n".join(fields)
                sections.append(content)
                records.append({"sheet": sheet.title, "row": row_number, "content": content})
                record_count += 1
            if header is None or record_count == 0:
                raise ValueError(f"Missing Excel reference table or records: {path}, {sheet.title}")
    finally:
        workbook.close()
    return "\n\n".join(sections), records


def parse_product_document(path: Path, effective_dates: dict | None = None) -> dict[str, object]:
    extension = path.suffix.lower()
    records = None
    if extension == ".md":
        text = path.read_text(encoding="utf-8")
        if not text.startswith("---\n"):
            raise ValueError(f"Missing YAML front matter: {path}")
        front_matter_parts = text.split("---\n", 2)
        if len(front_matter_parts) != 3:
            raise ValueError(f"Invalid YAML front matter: {path}")
        loaded_metadata = yaml.safe_load(front_matter_parts[1])
        if not isinstance(loaded_metadata, dict):
            raise ValueError(f"Invalid YAML metadata: {path}")
        metadata = dict(loaded_metadata)
        content = front_matter_parts[2].strip()
    else:
        if extension == ".pdf":
            pages = PyPDFLoader(str(path), mode="page").load()
            page_contents = []
            for page in pages:
                lines = []
                for line in page.page_content.splitlines():
                    # 去掉当前指南重复的页眉，页面正文连续合并后再按业务段落切分。
                    if not re.fullmatch(r"\S+\s*·\s*[\d.]+\s*·\s*\d+", line.strip()):
                        lines.append(line)
                page_contents.append("\n".join(lines))
            content = "\n".join(page_contents).strip()
        elif extension == ".docx":
            loaded_documents = Docx2txtLoader(str(path)).load()
            content = "\n\n".join(document.page_content for document in loaded_documents).strip()
        elif extension == ".xlsx":
            content, records = read_excel_document(path)
        else:
            raise ValueError(f"Unsupported product document: {path}")
        if effective_dates is None:
            effective_dates = product_effective_dates()
        metadata = extract_source_metadata(content, path, effective_dates)
    document = validate_document_metadata(metadata, path)
    if not content.strip():
        raise ValueError(f"Product document has no readable content: {path}")
    document["path"] = path
    document["content"] = content
    if records is not None:
        document["records"] = records
    return document


def product_effective_dates() -> dict:
    periods: dict = {}
    for path in discover_product_documents():
        if path.suffix.lower() != ".md":
            continue
        document = parse_product_document(path)
        scope = (document["company_id"], document["product"], document["version"])
        periods.setdefault(scope, set()).add((document["effective_from"], document["effective_to"]))
    return periods


def load_product_documents() -> list[dict[str, object]]:
    effective_dates = product_effective_dates()
    return [parse_product_document(path, effective_dates) for path in discover_product_documents()]


def build_document_chunks(document: dict[str, object]) -> list[dict[str, object]]:
    # 未增数据库列的业务元数据保存在每个 Chunk 正文中；现有范围与日期仍写入文档表。
    prefix = f"标题：{document['title']}\n公司：{document['company_id']} | 产品：{document['product']} | 版本：{document['version']} | 状态：{document['status']}\n最后核对日期：{document['last_reviewed']} | 主题：{document['topic']}\n\n"
    chunk_bodies: list[str] = []
    if "records" in document:
        for record in document["records"]:
            chunk_bodies.append(f"工作表：{record['sheet']} | 行：{record['row']}\n{record['content']}")
    else:
        content_size = CHUNK_SIZE - len(prefix)
        if content_size <= CHUNK_OVERLAP:
            raise ValueError(f"Document metadata is too long: {document['source_uri']}")
        splitter = RecursiveCharacterTextSplitter(chunk_size=content_size, chunk_overlap=CHUNK_OVERLAP, separators=["\n## ", "\n# ", "\n\n", "\n", "。", "！", "？", ". ", "；", ";", "，", ",", " ", ""])
        chunk_bodies = splitter.split_text(str(document["content"]))
    if not chunk_bodies:
        raise ValueError(f"Document produced no chunks: {document['source_uri']}")
    document_key = f"{document['company_id']}:{document['source_uri']}:{document['version']}"
    document_id = str(uuid5(NAMESPACE_URL, document_key))
    chunks: list[dict[str, object]] = []
    for index, body in enumerate(chunk_bodies):
        content = prefix + body
        hash_text = json.dumps([document_key, index, content], ensure_ascii=False)
        content_hash = hashlib.sha256(hash_text.encode("utf-8")).hexdigest()
        chunk_id = str(uuid5(NAMESPACE_URL, f"{document_key}:{index}:{content_hash}"))
        chunks.append({"chunk_id": chunk_id, "document_id": document_id, "company_id": document["company_id"], "content": content, "content_hash": content_hash})
    return chunks


@traceable(name="import_product_documents", run_type="chain")
def import_product_documents() -> dict[str, int]:
    documents = load_product_documents()
    all_chunks: list[dict[str, object]] = []
    for document in documents:
        chunks = build_document_chunks(document)
        document["chunks"] = chunks
        all_chunks.extend(chunks)
    with get_connection() as connection:
        rows = connection.execute("SELECT chunk_id::text, content_hash FROM support.document_chunks").fetchall()
    existing_hashes = {row["chunk_id"]: row["content_hash"] for row in rows}
    pending_chunks = []
    for chunk in all_chunks:
        if existing_hashes.get(chunk["chunk_id"]) != chunk["content_hash"]:
            pending_chunks.append(chunk)
    # 所有文件先读取、校验和向量化成功，再用一个事务同步知识表；失败保留原来的知识库。
    for offset in range(0, len(pending_chunks), EMBEDDING_BATCH_SIZE):
        batch = pending_chunks[offset:offset + EMBEDDING_BATCH_SIZE]
        for attempt in range(3):
            try:
                vectors = generate_text_embeddings([str(chunk["content"]) for chunk in batch], "RETRIEVAL_DOCUMENT")
                break
            except errors.ClientError as error:
                if error.code != 429 or attempt == 2:
                    raise
                delay = 60
                for detail in error.details.get("error", {}).get("details", []):
                    if detail.get("@type", "").endswith("RetryInfo"):
                        seconds = str(detail.get("retryDelay", ""))
                        if re.fullmatch(r"\d+(?:\.\d+)?s", seconds):
                            delay = min(60, max(1, math.ceil(float(seconds[:-1])) + 1))
                logger.warning("Embedding batch rate limited; retrying in %s seconds", delay)
                time.sleep(delay)
        for chunk, vector in zip(batch, vectors, strict=True):
            chunk["embedding"] = vector
    imported_count = 0
    removed_chunks = 0
    current_document_ids: list[str] = []
    source_prefix = PRODUCT_DOCUMENTS_DIR.relative_to(PROJECT_ROOT).as_posix() + "/"
    with get_connection() as connection:
        for document in documents:
            chunks = document["chunks"]
            row = connection.execute("""INSERT INTO support.product_documents (document_id, company_id, title, source_uri, product, version, effective_from, effective_to)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (company_id, source_uri, version) DO UPDATE SET title = EXCLUDED.title, product = EXCLUDED.product, effective_from = EXCLUDED.effective_from, effective_to = EXCLUDED.effective_to
                RETURNING document_id::text""", (chunks[0]["document_id"], document["company_id"], document["title"], document["source_uri"], document["product"], document["version"], document["effective_from"], document["effective_to"])).fetchone()
            document_id = row["document_id"]
            current_document_ids.append(document_id)
            chunk_ids = [chunk["chunk_id"] for chunk in chunks]
            removed_chunks += connection.execute("DELETE FROM support.document_chunks WHERE document_id = %s AND NOT (chunk_id = ANY(%s::uuid[]))", (document_id, chunk_ids)).rowcount
            changed = False
            for chunk in chunks:
                if "embedding" not in chunk:
                    continue
                connection.execute("""INSERT INTO support.document_chunks (chunk_id, document_id, company_id, content, content_hash, embedding)
                    VALUES (%s, %s, %s, %s, %s, %s::vector)
                    ON CONFLICT (chunk_id) DO UPDATE SET document_id = EXCLUDED.document_id, company_id = EXCLUDED.company_id, content = EXCLUDED.content, content_hash = EXCLUDED.content_hash, embedding = EXCLUDED.embedding""", (chunk["chunk_id"], document_id, chunk["company_id"], chunk["content"], chunk["content_hash"], json.dumps(chunk["embedding"])))
                changed = True
            if changed:
                imported_count += 1
        removed_chunks += connection.execute("""SELECT COUNT(*) AS count FROM support.document_chunks c JOIN support.product_documents d ON d.document_id = c.document_id
            WHERE starts_with(d.source_uri, %s) AND NOT (d.document_id = ANY(%s::uuid[]))""", (source_prefix, current_document_ids)).fetchone()["count"]
        removed_documents = connection.execute("DELETE FROM support.product_documents WHERE starts_with(source_uri, %s) AND NOT (document_id = ANY(%s::uuid[]))", (source_prefix, current_document_ids)).rowcount
    return {"found": len(documents), "imported": imported_count, "skipped": len(documents) - imported_count, "chunks": len(all_chunks), "embedded": len(pending_chunks), "reused": len(all_chunks) - len(pending_chunks), "removed_documents": removed_documents, "removed_chunks": removed_chunks}


# Product Markdown / PDF / DOCX / XLSX
# ↓
# parse_product_document()
# ↓
# metadata + 正文 / 完整表格记录
# ↓
# build_document_chunks() + 检查 content_hash
# ↓
# generate_text_embeddings()
# ↓
# 1024 维 Chunk Embedding
# ↓
# normalize
# ↓
# PostgreSQL / pgvector

# Document Loading / Splitting
# → 可以用 LangChain

# Embedding
# → 可以直接用官方 SDK

# Storage
# → 自己 PostgreSQL + pgvector

# Retrieval
# → 自己控制 SQL / Hybrid / Rerank

# Evaluation
# → 自己统一测试
